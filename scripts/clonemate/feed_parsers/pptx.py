"""PPTX parser — requires python-pptx (`pip install clonemate[pptx]`)."""
from __future__ import annotations

import datetime as _dt
import hashlib

import pptx  # python-pptx; ImportError → registry skips this parser

from clonemate.feed_parsers._base import FeedInput, _BaseParser
from clonemate.raw_writer import RawCandidate


class Parser(_BaseParser):
    name = "pptx"
    extras_hint = "pptx"

    def can_handle(self, fi: FeedInput) -> bool:
        return fi.path is not None and fi.path.suffix.lower() == ".pptx"

    def parse(self, fi: FeedInput) -> list[RawCandidate]:
        assert fi.path is not None
        prs = pptx.Presentation(str(fi.path))
        chunks: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            chunks.append(f"## Slide {i}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    chunks.append(shape.text)
        body = "\n\n".join(chunks)
        digest = hashlib.sha256(body.encode("utf-8")).hexdigest()[:12]
        now = (
            _dt.datetime.now(tz=_dt.timezone.utc)
            .astimezone()
            .isoformat(timespec="seconds")
        )
        return [
            RawCandidate(
                source_type="feed_pptx",
                relative_path=f"feed/pptx/{fi.path.stem}-{digest}.md",
                hash_input=body,
                frontmatter={
                    "src_type": "feed_pptx",
                    "src_origin": fi.path.name,
                    "collected_at": now,
                },
                content=body,
            )
        ]
